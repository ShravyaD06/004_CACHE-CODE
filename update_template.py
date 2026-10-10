import os
import sys
import pptx
from pptx.util import Pt

sys.stdout.reconfigure(encoding='utf-8')

template_path = r"c:\Users\devar\OneDrive\Desktop\DocuSense - AI\CBC2.0_Design_Round_Template.pptx"
output_dir = r"c:\Users\devar\OneDrive\Desktop\DocuSense - AI"

prs = pptx.Presentation(template_path)

def set_text(shape, new_text, font_size=None, bold=None):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    first_p = tf.paragraphs[0]
    first_p.text = new_text
    if font_size:
        first_p.font.size = Pt(font_size)
    if bold is not None:
        first_p.font.bold = bold
    while len(tf.paragraphs) > 1:
        p_elem = tf.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)

def set_text_multiline(shape, lines, font_size=13):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = line
        p.font.size = Pt(font_size)

# Update Master Layout Footers with Team ID: 004 and Domain: AI & ML
for master in prs.slide_masters:
    for layout in master.slide_layouts:
        for shape in layout.shapes:
            if shape.has_text_frame and ("Team ID:" in shape.text_frame.text or "Domain:" in shape.text_frame.text):
                shape.text_frame.text = "CBC 2.0   |   Design Round   |   Team ID: 004   |   Domain: AI & ML"

# =============================================================
# SLIDE 1 — TITLE
# =============================================================
slide1 = prs.slides[0]
set_text(slide1.shapes[2], "DocuSense AI  —  Extract. Validate. Verify. Correct.", font_size=26, bold=True)
set_text(slide1.shapes[5], "CACHE CODE", font_size=16, bold=True)
set_text(slide1.shapes[8], "004", font_size=16, bold=True)
set_text(slide1.shapes[11], "AI & ML", font_size=16, bold=True)
set_text(slide1.shapes[14], "Sapthagiri NPS University", font_size=16, bold=True)

# =============================================================
# SLIDE 2 — PROBLEM STATEMENT
# =============================================================
slide2 = prs.slides[1]
problem_text_exact = """Organizations receive invoices, application forms, certificates, and other documents in different layouts. Extracting information manually is repetitive and error-prone.

Objective: Build AI-powered system that extracts predefined information from uploaded documents and validates results.

Workflow:
- User uploads document or image.
- System identifies document type.
- Relevant fields extracted.
- Extracted values validated against expected formats/rules.
- Users review and correct uncertain fields.

Key Requirements:
- Support multiple document types.
- Extract configurable fields.
- Identify low-confidence or missing values.
- Validate dates, numbers, identifiers, and required fields.
- Allow users to correct extracted information.

Edge Cases:
- Poor-quality scan.
- Missing required field.
- Same information appears in multiple locations with different values.

Scope:
- Must Have: document upload, classification, field extraction, validation, review.
- Stretch: highlight source location of each extracted value.
- Demo: Upload different documents and demonstrate extraction, validation, correction of uncertain information."""

set_text(slide2.shapes[5], problem_text_exact, font_size=10)
set_text(slide2.shapes[8], "Problem Statement 3", font_size=13.5, bold=True)
set_text(slide2.shapes[11], "AI & ML", font_size=14, bold=True)
set_text(slide2.shapes[14], "Reliable extraction of configurable document fields, followed by validation and human review.", font_size=13, bold=True)

# =============================================================
# SLIDE 3 — UNDERSTANDING THE PROBLEM
# =============================================================
slide3 = prs.slides[2]
set_text_multiline(slide3.shapes[5], [
    "• Operations, finance, HR, admissions, and document-verification teams.",
    "• Staff handling invoices, certificates, and application forms.",
    "• Organizations processing documents in different layouts."
], font_size=13)

set_text_multiline(slide3.shapes[10], [
    "• Manual data entry is repetitive and error-prone.",
    "• Different document layouts make rigid extraction unreliable.",
    "• Missing, invalid, or conflicting values may be overlooked."
], font_size=13)

set_text_multiline(slide3.shapes[15], [
    "• Errors can delay document processing and verification.",
    "• Staff spend time on repetitive checking.",
    "• Reviewers need traceable evidence to trust extracted values."
], font_size=13)

# =============================================================
# SLIDE 4 — OUR PROPOSED SOLUTION
# =============================================================
slide4 = prs.slides[3]
set_text(slide4.shapes[5], "DocuSense AI helps document-processing teams turn varied documents into structured, validated information through OCR, configurable field extraction, and human review.", font_size=14, bold=True)

set_text(slide4.shapes[11], "Users upload supported PDFs, scanned PDFs, or JPG/PNG images.", font_size=13)
set_text(slide4.shapes[16], "The system reads text, identifies the likely document type, extracts configured fields, validates values, and flags missing, invalid, or conflicting information.", font_size=13)
set_text(slide4.shapes[21], "Users inspect flagged fields, correct uncertain results, and save reviewed structured information.", font_size=13)

# =============================================================
# SLIDE 5 — KEY FEATURES
# =============================================================
slide5 = prs.slides[4]
set_text(slide5.shapes[4], "Feature 1 — Multi-Format Upload", font_size=14, bold=True)
set_text(slide5.shapes[5], "Accept supported PDFs and images, and clearly reject unsupported or corrupted files.", font_size=12)

set_text(slide5.shapes[9], "Feature 2 — Document Classification", font_size=14, bold=True)
set_text(slide5.shapes[10], "Identify likely document categories such as invoices, certificates, and application forms.", font_size=12)

set_text(slide5.shapes[14], "Feature 3 — Configurable Field Extraction", font_size=14, bold=True)
set_text(slide5.shapes[15], "Use document-specific schemas to map relevant information into structured JSON.", font_size=12)

set_text(slide5.shapes[19], "Feature 4 — Validation Engine", font_size=14, bold=True)
set_text(slide5.shapes[20], "Check required fields, dates, numbers, identifiers, and configurable rules.", font_size=12)

set_text(slide5.shapes[24], "Feature 5 — Confidence-Aware Review", font_size=14, bold=True)
set_text(slide5.shapes[25], "Flag missing, uncertain, or conflicting values for human verification.", font_size=12)

set_text(slide5.shapes[29], "Feature 6 — Evidence-Linked Extraction", font_size=14, bold=True)
set_text(slide5.shapes[30], "Proposed stretch feature: show source text, page references, and bounding-box highlights when OCR coordinates are available.", font_size=11.5)

# =============================================================
# SLIDE 6 — ARCHITECTURE & TECH STACK
# =============================================================
slide6 = prs.slides[5]
set_text(slide6.shapes[5], "React + TypeScript + Tailwind CSS for web dashboard, upload, and review interface.", font_size=13)
set_text(slide6.shapes[10], "Python + FastAPI for document upload, processing APIs, validation, review, and result retrieval.", font_size=13)
set_text(slide6.shapes[15], "PyMuPDF & OpenCV for parsing; PaddleOCR + optional LLM API; SQLite DB for prototype.", font_size=13)

set_text(slide6.shapes[19], "Figma", font_size=13)
set_text(slide6.shapes[21], "React / TypeScript", font_size=13)
set_text(slide6.shapes[23], "Python / FastAPI", font_size=13)
set_text(slide6.shapes[25], "SQLite", font_size=13)
set_text(slide6.shapes[27], "PaddleOCR / LLM API", font_size=13)

# =============================================================
# SLIDE 7 — INNOVATION & DIFFERENTIATION
# =============================================================
slide7 = prs.slides[6]
table = slide7.shapes[1].table

table.cell(1, 1).text = "Manual or rigid data-entry workflows."
table.cell(1, 2).text = "One upload-and-review workflow."

table.cell(2, 1).text = "Repetitive staff effort and tooling costs."
table.cell(2, 2).text = "Accessible software stack for prototype; API limits apply."

table.cell(3, 1).text = "Layout-specific extraction rules required."
table.cell(3, 2).text = "Configurable field schemas for different categories."

table.cell(4, 1).text = "Synchronous workflows become bottlenecks."
table.cell(4, 2).text = "Asynchronous jobs & processing workers proposed for multiple uploads."

for row in table.rows:
    for cell in row.cells:
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11.5)

set_text(slide7.shapes[3], "Our unique edge: DocuSense AI combines configurable field extraction, validation, conflict detection, and evidence-linked human review in one workflow.", font_size=13.5, bold=True)

# =============================================================
# SLIDE 8 — IMPACT, FEASIBILITY & RISKS
# =============================================================
slide8 = prs.slides[7]
set_text(slide8.shapes[6], "Reduce repetitive data-entry work and help reviewers detect missing, invalid, and conflicting values. Actual time savings must be measured during testing.", font_size=12)

set_text(slide8.shapes[12], "A four-person team can target an end-to-end prototype using React, FastAPI, PaddleOCR, Python validation rules, and SQLite within the 18-hour implementation window.", font_size=12)

set_text(slide8.shapes[18], "Use background jobs, job-status tracking, and worker processes. PostgreSQL, object storage, and a distributed queue are possible future upgrades.", font_size=12)

set_text_multiline(slide8.shapes[24], [
    "• Poor scans: preprocessing & quality warnings.",
    "• Missing fields: explicit required-field flags.",
    "• Conflicting values: retain alternatives & show source evidence.",
    "• API failure: timeouts, safe retries & fallback.",
    "• Sensitive documents: access controls & cautious API use.",
    "• Large uploads: queued processing & performance testing."
], font_size=10.5)

# =============================================================
# SLIDE 9 — RULES TO FOLLOW (PRESERVE EXACTLY)
# =============================================================
# Slide 9 rules preserved untouched.

# Target Output File Paths
target_pptx_paths = [
    os.path.join(output_dir, "004_AI_ML.pptx"),
    os.path.join(output_dir, "TeamID_Domain.pptx"),
    os.path.join(output_dir, "CACHE_CODE_AI_ML.pptx"),
    os.path.join(output_dir, "CBC2.0_Design_Round_Template_Completed.pptx")
]

for path in target_pptx_paths:
    prs.save(path)
    print(f"Saved PPTX: {path}")

# Export PDF versions
import win32com.client

ppt_app = win32com.client.Dispatch("PowerPoint.Application")
ppt_app.Visible = True

try:
    for pptx_path in target_pptx_paths:
        pdf_path = os.path.splitext(pptx_path)[0] + ".pdf"
        abs_pptx = os.path.abspath(pptx_path)
        abs_pdf = os.path.abspath(pdf_path)
        
        presentation = ppt_app.Presentations.Open(abs_pptx, WithWindow=False)
        presentation.SaveAs(abs_pdf, 32)
        presentation.Close()
        print(f"Exported PDF: {abs_pdf}")
finally:
    ppt_app.Quit()

print("All PPTX and PDF files generated and saved successfully!")
